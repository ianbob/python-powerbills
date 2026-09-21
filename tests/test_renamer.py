import tempfile
import unittest
from pathlib import Path
from typing import cast

from powerbills.renamer import (
    extract_account_number,
    extract_bill_data,
    load_station_lookup,
    merge_bill_data,
    process_files,
)


class ExtractAccountNumberTests(unittest.TestCase):
    def test_extracts_labelled_number_with_spacing(self):
        self.assertEqual(extract_account_number("Account Number: 123 456 789"), "123456789")

    def test_uses_account_row_when_ocr_drops_punctuation(self):
        self.assertEqual(extract_account_number("ACCOUNT 123456789"), "123456789")

    def test_ignores_unrelated_numbers(self):
        self.assertIsNone(extract_account_number("Invoice 123456789\nMeter 456789"))

    def test_extracts_requested_bill_fields(self):
        bill = extract_bill_data(
            "ACCOUNT NUMBER: 175565209\n"
            "Date Due: 09/09/2026\n"
            "Consumption Period: 23/07/2026-22/08/2026\n"
            "ELECTRICITY BILL: Ksh 103.00\n"
            "HighRateConsumption 3kWh x 12.28\n"
            "LowRateConsumption 2kWh x 12.28"
        )
        assert bill is not None
        self.assertEqual(bill.account_number, "175565209")
        self.assertEqual(bill.bill, 103.0)
        self.assertEqual(bill.due_date, "09/09/2026")
        self.assertEqual(bill.low_rate_consumption, 2.0)
        self.assertEqual(bill.high_rate_consumption, 3.0)


class ProcessFilesTests(unittest.TestCase):
    def test_preview_does_not_rename(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "bill.pdf"
            source.write_bytes(b"pdf")
            result = process_files([source], ocr=lambda _: "Account Number: 123456789")[0]
            self.assertEqual(result.status, "planned")
            self.assertTrue(source.exists())

    def test_apply_renames_file(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "bill.pdf"
            source.write_bytes(b"pdf")
            result = process_files(
                [source], apply=True, ocr=lambda _: "Account Number: 123456789"
            )[0]
            self.assertEqual(result.status, "renamed")
            self.assertTrue((Path(folder) / "123456789.pdf").exists())

    def test_account_header_wins_over_other_numbers(self):
        text = "ACCOUNT NUMBER: 178761435\nInvoice Number: 260910015872316\nMeter Number: 040016111611"
        self.assertEqual(process_files([Path("bill.pdf")], ocr=lambda _: text)[0].message, "178761435")

    def test_uses_station_name_from_lookup(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "178761435.pdf"
            source.write_bytes(b"pdf")
            result = process_files(
                [source],
                apply=True,
                ocr=lambda _: "Account Number: 178761435",
                station_lookup={"178761435": "Athi River"},
            )[0]
            self.assertEqual(result.message, "178761435")
            self.assertTrue((Path(folder) / "Athi River.pdf").exists())

    def test_falls_back_to_account_when_station_is_missing(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "bill.pdf"
            source.write_bytes(b"pdf")
            process_files(
                [source],
                apply=True,
                ocr=lambda _: "Account Number: 999999999",
                station_lookup={"178761435": "Athi River"},
            )
            self.assertTrue((Path(folder) / "999999999.pdf").exists())


class WorkbookLookupTests(unittest.TestCase):
    def test_loads_requested_columns(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            workbook_path = Path(folder) / "lookup.xlsx"
            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "august"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", 178924769])
            workbook.save(workbook_path)
            self.assertEqual(
                load_station_lookup(workbook_path), {"178924769": "Athi River"}
            )

    def test_loads_the_only_sheet_when_its_name_changes(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            workbook_path = Path(folder) / "lookup.xlsx"
            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "September"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", 178924769])
            workbook.save(workbook_path)
            self.assertEqual(
                load_station_lookup(workbook_path), {"178924769": "Athi River"}
            )

    def test_merges_bill_fields_into_template(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / "template.xlsx"
            output = Path(folder) / "output.xlsx"
            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "Sheet1"
            sheet.append(["Station", "Meter Number", "Account Number"])
            sheet.append(["Athi River", "123", 178924769])
            workbook.save(template)
            bill = extract_bill_data(
                "ACCOUNT NUMBER: 178924769\nDate Due: 09/09/2026\n"
                "Consumption Period: 23/07/2026-22/08/2026\n"
                "ELECTRICITY BILL: Ksh 103.00"
            )
            assert bill is not None
            merge_bill_data(
                template,
                output,
                [bill],
            )
            result = load_workbook(output, data_only=True)["Sheet1"]
            headers = [cell.value for cell in result[1]]
            values = [cell.value for cell in result[2]]
            expected = {
                "Consumption Period": "23/07/2026-22/08/2026",
                "Bill": 103,
                "Due Date": "09/09/2026",
                "LowRateConsumption": None,
                "HighRateConsumption": None,
            }
            for header, expected_value in expected.items():
                self.assertEqual(values[headers.index(header)], expected_value)

class ProcessFilesWithLookupTests(unittest.TestCase):
    def test_uses_station_name_from_lookup_file(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "178761435.pdf"
            source.write_bytes(b"pdf")
            lookup_file = Path(folder) / "lookup.xlsx"
            from openpyxl import Workbook

            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "Sheet1"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", 178761435])
            workbook.save(lookup_file)
            result = process_files(
                [source],
                apply=True,
                ocr=lambda _: "Account Number: 178761435",
                station_lookup=load_station_lookup(lookup_file),
            )[0]
            self.assertEqual(result.message, "178761435")
            self.assertTrue((Path(folder) / "Athi River.pdf").exists())

class CliDefaultsTests(unittest.TestCase):
    def test_production_input_directory_is_raw_files(self):
        from powerbills.cli import DEFAULT_INPUT_DIRECTORY

        self.assertEqual(DEFAULT_INPUT_DIRECTORY, "raw_files")


if __name__ == "__main__":
    unittest.main()